"""Importacao de pacientes: leitura de CSV/XLSX, validacao (CPF, telefone, data, e-mail), duplicados e confirmacao."""

import csv
import io
import unicodedata
from datetime import date, datetime

from django.conf import settings
from django.core.exceptions import ValidationError
from django.core.validators import validate_email
from django.db import transaction

from apps.auditoria import servico as auditoria
from apps.core.validadores import cpf_valido, somente_digitos
from apps.pacientes.models import Paciente
from apps.pacientes.servico import profissional_da_requisicao, vincular_profissional

from .models import LoteImportacao

COLUNAS = ["nome", "cpf", "nascimento", "email", "telefone", "convenio", "carteirinha"]
MODELO_CSV = (
    "﻿nome;cpf;nascimento;email;telefone;convenio;carteirinha\r\n"
    "Maria da Silva;529.982.247-25;15/03/1985;maria@exemplo.com;(11) 98888-7777;Unimed;0012345\r\n"
    "João Souza;;02/11/1990;;11977776666;;\r\n"
)
ALIAS = {
    "nome": "nome", "nomecompleto": "nome", "paciente": "nome",
    "cpf": "cpf",
    "nascimento": "nascimento", "datadenascimento": "nascimento", "datanascimento": "nascimento", "dtnascimento": "nascimento",
    "email": "email", "emailcontato": "email",
    "telefone": "telefone", "celular": "telefone", "fone": "telefone", "whatsapp": "telefone",
    "convenio": "convenio", "plano": "convenio", "operadora": "convenio",
    "carteirinha": "carteirinha", "carteira": "carteirinha", "matricula": "carteirinha",
}


class ErroImportacao(Exception):
    """Arquivo ou regra invalida; a mensagem pode ser mostrada ao usuario."""


def _chave(texto) -> str:
    sem_acento = unicodedata.normalize("NFKD", str(texto)).encode("ascii", "ignore").decode()
    return "".join(c for c in sem_acento.lower() if c.isalnum())


def _mapear_cabecalho(cabecalho):
    mapa = {}
    for posicao, titulo in enumerate(cabecalho):
        coluna = ALIAS.get(_chave(titulo))
        if coluna and coluna not in mapa.values():
            mapa[posicao] = coluna
    if "nome" not in mapa.values():
        raise ErroImportacao("Não encontramos a coluna “nome”. Use o modelo de planilha.")
    return mapa


# --------------------------------------------------------------------------- leitura


def _ler_csv(conteudo: bytes):
    try:
        texto = conteudo.decode("utf-8-sig")
    except UnicodeDecodeError:
        texto = conteudo.decode("latin-1")
    amostra = texto[:2000]
    delimitador = ";" if amostra.count(";") >= amostra.count(",") else ","
    return list(csv.reader(io.StringIO(texto), delimiter=delimitador))


def _ler_xlsx(conteudo: bytes):
    from openpyxl import load_workbook

    try:
        planilha = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
    except Exception as exc:
        raise ErroImportacao("Não foi possível abrir a planilha. Salve como .xlsx ou .csv e tente de novo.") from exc
    return [list(linha) for linha in planilha.active.iter_rows(values_only=True)]


def ler_arquivo(arquivo):
    nome = (arquivo.name or "").lower()
    if arquivo.size > settings.MEUPSIQ_IMPORTACAO_MAX_BYTES:
        raise ErroImportacao(f"O arquivo passa do limite de {settings.MEUPSIQ_IMPORTACAO_MAX_BYTES // (1024 * 1024)} MB.")
    conteudo = arquivo.read()
    if nome.endswith(".csv") or nome.endswith(".txt"):
        linhas = _ler_csv(conteudo)
    elif nome.endswith(".xlsx"):
        linhas = _ler_xlsx(conteudo)
    else:
        raise ErroImportacao("Use um arquivo .csv ou .xlsx.")
    linhas = [l for l in linhas if any(str(c).strip() for c in l if c is not None)]
    if len(linhas) < 2:
        raise ErroImportacao("A planilha está vazia. Preencha ao menos uma linha abaixo do cabeçalho.")
    if len(linhas) - 1 > settings.MEUPSIQ_IMPORTACAO_MAX_LINHAS:
        raise ErroImportacao(f"A planilha tem mais de {settings.MEUPSIQ_IMPORTACAO_MAX_LINHAS} linhas. Divida em partes.")
    mapa = _mapear_cabecalho(linhas[0])
    resultado = []
    for numero, linha in enumerate(linhas[1:], start=2):
        resultado.append((numero, {coluna: (linha[pos] if pos < len(linha) else None) for pos, coluna in mapa.items()}))
    return resultado


# --------------------------------------------------------------------------- validacao


def _texto(valor, limite):
    if valor is None:
        return ""
    return " ".join(str(valor).split())[:limite]


def _data(valor):
    if valor in (None, ""):
        return None
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    texto = str(valor).strip()
    for formato in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%d/%m/%y"):
        try:
            return datetime.strptime(texto, formato).date()
        except ValueError:
            continue
    raise ValueError


def validar_linha(numero, bruto, cpfs_existentes, nomes_nascimentos, cpfs_do_arquivo):
    msgs, avisos, status = [], [], "ok"
    nome = _texto(bruto.get("nome"), 160)
    if not nome:
        msgs.append("Nome vazio.")

    cpf_bruto = bruto.get("cpf")
    cpf = somente_digitos(str(cpf_bruto)) if cpf_bruto is not None else ""
    if isinstance(cpf_bruto, (int, float)) and cpf and len(cpf) < 11:  # Excel remove zeros a esquerda
        cpf = cpf.zfill(11)
    if cpf and not cpf_valido(cpf):
        msgs.append("CPF inválido.")
        cpf = ""

    nascimento = None
    try:
        nascimento = _data(bruto.get("nascimento"))
        if nascimento and nascimento > date.today():
            msgs.append("Data de nascimento no futuro.")
            nascimento = None
    except ValueError:
        msgs.append("Data de nascimento inválida (use dd/mm/aaaa).")

    email = _texto(bruto.get("email"), 254).lower()
    if email:
        try:
            validate_email(email)
        except ValidationError:
            msgs.append("E-mail inválido.")
            email = ""

    telefone_bruto = bruto.get("telefone")
    telefone = somente_digitos(str(int(telefone_bruto)) if isinstance(telefone_bruto, float) else str(telefone_bruto or ""))
    if telefone and not 10 <= len(telefone) <= 13:
        msgs.append("Telefone inválido (informe DDD e número).")
        telefone = ""

    if msgs:
        status = "erro"
    elif cpf and cpf in cpfs_existentes:
        status, msgs = "duplicado", [f"CPF já cadastrado para {cpfs_existentes[cpf]}."]
    elif cpf and cpf in cpfs_do_arquivo:
        status, msgs = "duplicado", ["CPF repetido na planilha (linha anterior)."]
    else:
        chave = (nome.casefold(), nascimento)
        if nascimento and chave in nomes_nascimentos:
            avisos.append("Mesmo nome e nascimento de um paciente já cadastrado. Confira antes de confirmar.")
    if status == "ok" and cpf:
        cpfs_do_arquivo.add(cpf)

    return {
        "linha": numero, "status": status, "mensagens": msgs, "avisos": avisos,
        "dados": {"nome": nome, "cpf": cpf, "nascimento": nascimento.isoformat() if nascimento else "", "email": email, "telefone": telefone,
                  "convenio": _texto(bruto.get("convenio"), 80), "carteirinha": _texto(bruto.get("carteirinha"), 40)},
    }


def criar_lote(request, arquivo) -> LoteImportacao:
    brutos = ler_arquivo(arquivo)
    consultorio = request.consultorio
    existentes = Paciente.objects.filter(consultorio=consultorio, mesclado_em__isnull=True)
    cpfs = {p.cpf: p.nome for p in existentes if p.cpf}
    nomes = {(p.nome.casefold(), p.nascimento) for p in existentes if p.nascimento}
    no_arquivo = set()
    linhas = [validar_linha(n, bruto, cpfs, nomes, no_arquivo) for n, bruto in brutos]
    lote = LoteImportacao.objects.create(
        consultorio=consultorio, criado_por_id=request.user.pk, nome_arquivo=(arquivo.name or "planilha")[:200], linhas=linhas,
        total=len(linhas), validas=sum(l["status"] == "ok" for l in linhas), duplicadas=sum(l["status"] == "duplicado" for l in linhas),
        com_erro=sum(l["status"] == "erro" for l in linhas),
    )
    auditoria.registrar(request, "importacao_previa", "lote_importacao", lote.pk, total=lote.total, validas=lote.validas)
    return lote


@transaction.atomic
def confirmar(request, lote) -> int:
    if lote.status != LoteImportacao.Status.PREVIA:
        raise ErroImportacao("Esta importação já foi concluída ou cancelada.")
    consultorio = request.consultorio
    cpfs_agora = set(Paciente.objects.filter(consultorio=consultorio, mesclado_em__isnull=True).exclude(cpf="").values_list("cpf", flat=True))
    profissional = profissional_da_requisicao(request)
    criados = 0
    for linha in lote.linhas:
        if linha["status"] != "ok":
            continue
        d = linha["dados"]
        if d["cpf"] and d["cpf"] in cpfs_agora:  # alguem cadastrou o mesmo CPF depois da previa
            linha["status"], linha["mensagens"] = "duplicado", ["CPF cadastrado depois da prévia."]
            continue
        paciente = Paciente.objects.create(
            consultorio=consultorio, nome=d["nome"], cpf=d["cpf"], nascimento=date.fromisoformat(d["nascimento"]) if d["nascimento"] else None,
            email=d["email"], telefone=d["telefone"], convenio=d["convenio"], carteirinha=d["carteirinha"],
        )
        if profissional:
            vincular_profissional(paciente, profissional)
        if d["cpf"]:
            cpfs_agora.add(d["cpf"])
        criados += 1
    lote.status, lote.importadas = LoteImportacao.Status.CONCLUIDA, criados
    lote.duplicadas = sum(l["status"] == "duplicado" for l in lote.linhas)
    lote.save()
    auditoria.registrar(request, "importacao_concluida", "lote_importacao", lote.pk, importados=criados)
    return criados


def cancelar(request, lote):
    if lote.status != LoteImportacao.Status.PREVIA:
        raise ErroImportacao("Esta importação já foi concluída ou cancelada.")
    lote.status, lote.linhas = LoteImportacao.Status.CANCELADA, []  # nao guarda os dados de uma importacao descartada
    lote.save()
    auditoria.registrar(request, "importacao_cancelada", "lote_importacao", lote.pk)
